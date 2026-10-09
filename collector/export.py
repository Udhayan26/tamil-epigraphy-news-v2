"""Export verified articles to Excel (.xlsx) and CSV."""
from __future__ import annotations

import csv
import json
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Config
from .db import DB

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")

COLUMNS = [
    ("No", 6), ("Date", 12), ("Year", 7), ("Title", 60), ("Keyword group (Tamil)", 22),
    ("Keyword group (English)", 26), ("Matched terms", 30), ("Places mentioned", 24),
    ("Source", 24), ("Source country", 14), ("Language", 12), ("URL", 55),
    ("Archive copy", 40), ("Snippet", 70), ("Match basis", 12), ("Found via", 22),
    ("Found on", 12),
]
FOUND_COL = len(COLUMNS)          # 1-based index of "Found on"
NEW_DAYS = 7


def _rows(db: DB, cfg: Config) -> list[list]:
    labels = {g.key: (g.label_ta, g.label_en) for g in cfg.groups}
    cur = db.conn.execute(
        "SELECT * FROM articles ORDER BY COALESCE(NULLIF(REPLACE(pub_date,'≤',''),''),'0000') DESC, title"
    )
    out = []
    for i, a in enumerate(cur.fetchall(), 1):
        groups = json.loads(a["groups"] or "[]")
        out.append([
            i, a["pub_date"] or "", a["year"] or "", a["title"] or "",
            ", ".join(labels.get(g, (g, g))[0] for g in groups),
            ", ".join(labels.get(g, (g, g))[1] for g in groups),
            ", ".join(json.loads(a["matched_terms"] or "[]")),
            ", ".join(json.loads(a["places"] or "[]")),
            a["domain"] or "", a["source_country"] or "", a["language"] or "",
            a["url"], a["archive_url"] or "", a["snippet"] or "",
            a["match_basis"] or "", a["collector"] or "", (a["fetched_at"] or "")[:10],
        ])
    return out


def export(db: DB, cfg: Config, out_dir: Path) -> dict:
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = _rows(db, cfg)
    header = [c for c, _ in COLUMNS]

    with open(out_dir / "articles.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)

    # Articles first found in the last NEW_DAYS days, and today's list as its own CSV.
    today = time.strftime("%Y-%m-%d", time.gmtime())
    cutoff = time.strftime("%Y-%m-%d", time.gmtime(time.time() - NEW_DAYS * 86400))
    recent = [r for r in rows if r[FOUND_COL - 1] >= cutoff]
    todays = [r for r in rows if r[FOUND_COL - 1] == today]
    if todays:
        daily_dir = out_dir / "daily"
        daily_dir.mkdir(parents=True, exist_ok=True)
        with open(daily_dir / f"{today}.csv", "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(todays)

    wb = Workbook()
    ws = wb.active
    ws.title = "Articles"
    _fill_sheet(ws, header, rows)
    last = max(ws.max_row, 2)
    _fill_sheet(wb.create_sheet(f"New (last {NEW_DAYS} days)"), header, recent)

    # ---- Summary sheets (formulas, so they update if you edit the Articles sheet) ----
    n = last
    s = wb.create_sheet("By keyword")
    s.append(["Keyword group (Tamil)", "Keyword group (English)", "Articles"])
    for g in cfg.groups:
        r = s.max_row + 1
        s.append([g.label_ta, g.label_en, f'=COUNTIF(Articles!$F$2:$F${n},"*"&B{r}&"*")'])
    s.append(["Total articles", "", f"=COUNTA(Articles!$A$2:$A${n})"])
    _style_summary(s, [24, 30, 12])

    y = wb.create_sheet("By year")
    y.append(["Year", "Articles"])
    years = sorted({r[2] for r in rows if r[2]}, reverse=True)
    for yr in years:
        rr = y.max_row + 1
        y.append([yr, f"=COUNTIF(Articles!$C$2:$C${n},A{rr})"])
    y.append(["Unknown year", f'=COUNTBLANK(Articles!$C$2:$C${n})'])
    _style_summary(y, [14, 12])

    src = wb.create_sheet("By source")
    src.append(["Source", "Source country", "Articles"])
    seen = {}
    for r in rows:
        seen.setdefault(r[8], r[9])
    for dom, ctry in sorted(seen.items()):
        rr = src.max_row + 1
        src.append([dom, ctry, f"=COUNTIF(Articles!$I$2:$I${n},A{rr})"])
    _style_summary(src, [32, 16, 12])

    pl = wb.create_sheet("By place")
    pl.append(["Place mentioned", "Articles"])
    for p in cfg.places:
        rr = pl.max_row + 1
        pl.append([p["name"], f'=COUNTIF(Articles!$H$2:$H${n},"*"&A{rr}&"*")'])
    _style_summary(pl, [24, 12])

    wb.calculation.fullCalcOnLoad = True
    wb.save(out_dir / "articles.xlsx")
    return {"articles": len(rows), "new_today": len(todays), "new_week": len(recent)}


def _fill_sheet(ws, header, rows) -> None:
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    url_col = header.index("URL") + 1
    arch_col = header.index("Archive copy") + 1
    for r in rows:
        ws.append(r)
        rn = ws.max_row
        for c in ws[rn]:
            c.font = BODY_FONT
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (4, 14))
        for col in (url_col, arch_col):
            cell = ws.cell(row=rn, column=col)
            if cell.value and str(cell.value).startswith("http") and len(str(cell.value)) < 2000:
                cell.hyperlink = str(cell.value)
                cell.font = LINK_FONT
    for i, (_, wdt) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 2)}"


def _style_summary(ws, widths):
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
