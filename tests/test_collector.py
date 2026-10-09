"""Offline tests: no network needed (HTTP calls are replaced with fakes)."""
import json
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from collector import http
from collector.config import load_config
from collector.db import DB
from collector.export import export
from collector.match import Matcher, is_tamil
from collector.sources import date_archive, gdelt, wayback
from collector.sources.util import date_from_url
from collector import verify


@pytest.fixture
def cfg():
    return load_config()


@pytest.fixture
def matcher(cfg):
    return Matcher(cfg)


@pytest.fixture
def db(tmp_path):
    return DB(tmp_path / "t.sqlite")


def resp(text="", status=200, url="https://x", ctype="text/html"):
    return SimpleNamespace(text=text, status_code=status, url=url, content=text.encode(),
                           headers={"content-type": ctype})


ARTICLE_EN = """<html><head><title>Hero stone found near Madurai</title>
<meta property="article:published_time" content="2020-10-14T10:00:00+05:30"></head>
<body><article><h1>Hero stone found near Madurai</h1>
<p>A 400-year-old hero stone was discovered at Munisalai in Madurai, Tamil Nadu, by
researchers studying temple inscriptions. The stone belongs to the Nayak period and shows a
warrior with a raised sword beside his wife.</p>
<p>Archaeologists said this is the first such find in the heart of the city.</p></article></body></html>"""

ARTICLE_TA = """<html><head><title>திருப்பூரில் தமிழ்-பிராமி கல்வெட்டு கண்டுபிடிப்பு</title></head>
<body><article><h1>திருப்பூரில் தமிழ்-பிராமி கல்வெட்டு கண்டுபிடிப்பு</h1>
<p>திருப்பூர் மாவட்டம் குமாரிக்கல் பாளையத்தில் இந்திய தொல்பொருள் ஆய்வுத் துறை மேற்கொண்ட ஆய்வில்,
தமிழ்-பிராமி எழுத்துக்கள் பொறிக்கப்பட்ட பானை ஓடு கண்டெடுக்கப்பட்டுள்ளது. இங்குள்ள 26 அடி உயர நடுகல்
பாதுகாக்கப்பட்ட நினைவுச் சின்னமாக அறிவிக்கப் படலாம். இது சங்க கால சேர மன்னர்களுடன் தொடர்புடையது.</p>
</article></body></html>"""

ARTICLE_OFFTOPIC = """<html><head><title>Roman villa excavated in Kent</title></head><body><article>
<h1>Roman villa excavated in Kent</h1><p>Archaeologists in England have excavated a Roman villa
with mosaic floors near Canterbury. The excavation revealed coins and pottery from the third
century and the team will continue digging next summer with volunteers from the county.</p>
</article></body></html>"""


# ---------------- matching ----------------
def test_english_match(matcher):
    r = matcher.accept("Hero stone found near Madurai", "A hero stone was found in Tamil Nadu.")
    assert r and "hero_stone" in r.groups and "Tamil Nadu" in r.places


def test_tamil_suffix_match(matcher):
    r = matcher.match("நடுகற்களை கண்டெடுத்தனர்")
    assert "hero_stone" in r.groups


def test_word_boundary(matcher):
    # 'erode' must not match 'eroded'; 'inscription' must not match 'subscription'
    r = matcher.match("Subscription prices eroded")
    assert not r.ok


def test_context_required(matcher):
    assert matcher.accept("Roman villa excavated", "Excavation in Kent, England.") is None
    assert matcher.accept("Chola inscription found", "near Thanjavur") is not None


def test_hyphen_variants(matcher):
    assert "inscription" in matcher.match("Tamil Brahmi script on pot").groups
    assert "inscription" in matcher.match("tamil-brahmi").groups


def test_is_tamil():
    assert is_tamil("நடுகல் கண்டுபிடிப்பு")
    assert not is_tamil("Hero stone found")


# ---------------- url dates ----------------
@pytest.mark.parametrize("url,expected", [
    ("http://www.hindu.com/2010/04/14/stories/2010041456210700.htm", "2010-04-14"),
    ("https://www.newindianexpress.com/states/tamil-nadu/2020/oct/14/hero-stone-123.html", "2020-10-14"),
    ("https://www.dtnext.in/tamilnadu/2017/04/02/hero-stones-of-women", "2017-04-02"),
    ("https://www.thehindu.com/news/x/article33679561.ece", None),
])
def test_date_from_url(url, expected):
    assert date_from_url(url) == expected


# ---------------- wayback ----------------
def test_wayback_plan_batches(cfg):
    jobs = wayback.plan(cfg)
    assert jobs and all(len(j["regex"]) <= wayback.MAX_REGEX_LEN + 200 for j in jobs)
    ta_jobs = [j for j in jobs if j["domain"] == "dinamani.com"]
    assert any("%E0%AE" in j["regex"] for j in ta_jobs)     # Tamil slugs included
    en_jobs = [j for j in jobs if j["domain"] == "thehindu.com"]
    assert not any("%E0%AE" in j["regex"] for j in en_jobs)


def test_wayback_step(monkeypatch, cfg, db):
    calls = []
    pages = [
        [["timestamp", "original"],
         ["20101014000000", "http://www.hindu.com/2010/10/14/stories/hero-stone-found.htm"],
         ["20101014000000", "http://www.hindu.com/photos/hero-stone.jpg"],
         [], ["RESUMEKEY1"]],
        [["timestamp", "original"],
         ["20110101000000", "http://www.hindu.com/2011/01/01/stories/inscription-found.htm"]],
    ]

    def fake_get(url, params=None, **kw):
        calls.append(dict(params))
        return resp(json.dumps(pages[len(calls) - 1]))

    monkeypatch.setattr(http, "get", fake_get)
    assert wayback.step(db, cfg)                       # first block, has resume key
    st = db.get_state(wayback.STATE)
    assert st == {"job": 0, "resume": "RESUMEKEY1"}
    assert calls[0]["showResumeKey"] == "true" and "page" not in calls[0]
    assert any(f.startswith("original:(?i)") for f in calls[0]["filter"])
    assert wayback.step(db, cfg)                       # second block, no key -> next job
    assert calls[1]["resumeKey"] == "RESUMEKEY1"
    assert db.get_state(wayback.STATE)["job"] == 1
    c = db.conn.execute("SELECT * FROM candidates ORDER BY url").fetchall()
    assert len(c) == 2 and c[0]["hint_date"] == "2010-10-14" and c[0]["url"].startswith("https://")
    assert "links found so far 2" in wayback.progress(db, cfg)


def test_wayback_errors_skip_job(monkeypatch, cfg, db):
    monkeypatch.setattr(http, "get", lambda *a, **k: resp("", status=503))
    monkeypatch.setattr(wayback.time, "sleep", lambda s: None)
    for _ in range(3):
        wayback.step(db, cfg)
    assert db.get_state(wayback.STATE)["job"] == 1
    assert len(db.get_state("wayback_skipped")) == 1


@pytest.mark.parametrize("text,n,key", [
    ("", 0, None),
    ("[]", 0, None),
    ('[["timestamp","original"],["2020","http://a"],[],["K"]]', 1, "K"),
    ('[["2020","http://a"],["2021","http://b"]]', 2, None),
])
def test_parse_cdx(text, n, key):
    rows, resume = wayback.parse_cdx(text)
    assert len(rows) == n and resume == key


# ---------------- gdelt ----------------
def test_gdelt_parse(monkeypatch):
    data = {"articles": [{"url": "https://www.thehindu.com/a.ece", "title": "Hero stone",
                          "seendate": "20210301T101500Z"}]}
    monkeypatch.setattr(http, "get", lambda *a, **k: resp(json.dumps(data)))
    from datetime import datetime
    rows = gdelt.search("x", datetime(2021, 3, 1), datetime(2021, 3, 31))
    assert rows == [{"url": "https://www.thehindu.com/a.ece", "title": "Hero stone",
                     "date": "2021-03-01", "collector": "gdelt"}]


def test_gdelt_text_error(monkeypatch):
    monkeypatch.setattr(http, "get", lambda *a, **k: resp("The specified phrase is too short."))
    from datetime import datetime
    assert gdelt.search("x", datetime(2021, 3, 1), datetime(2021, 3, 31)) == []


# ---------------- date archive ----------------
def test_date_archive(monkeypatch, cfg, db, matcher):
    page = """<html><body>
      <a href="https://www.thehindu.com/2010/04/14/stories/1.htm">Inscription of Aditya Chola I discovered</a>
      <a href="https://www.thehindu.com/2010/04/14/stories/2.htm">Cricket team wins the series again</a>
    </body></html>"""
    monkeypatch.setattr(http, "get", lambda *a, **k: resp(page))
    assert date_archive.step(db, cfg, matcher)
    c = db.conn.execute("SELECT * FROM candidates").fetchall()
    assert len(c) == 1 and "Aditya" in c[0]["hint_title"]


# ---------------- verify + export ----------------
def test_verify_and_export(monkeypatch, cfg, db, matcher, tmp_path):
    pages = {
        "https://www.etvbharat.com/hero": ARTICLE_EN,
        "https://www.dinamani.com/ta": ARTICLE_TA,
        "https://www.heritagedaily.com/kent": ARTICLE_OFFTOPIC,
    }

    def fake_get(url, **kw):
        if url in pages:
            return resp(pages[url], url=url)
        return resp("", status=404, url=url)

    monkeypatch.setattr(http, "get", fake_get)
    for u in pages:
        db.add_candidate(u, "wayback:test")
    db.add_candidate("https://dead.example.lk/x", "wayback:test")
    db.add_candidate("https://news.google.com/rss/articles/abc", "gnews",
                     hint_title="Hero stone unearthed in Krishnagiri", hint_date="2026-10-01")
    db.commit()

    verify.verify_batch(db, cfg, matcher, workers=2)
    arts = {a["url"]: a for a in db.conn.execute("SELECT * FROM articles")}
    assert set(arts) == {"https://www.etvbharat.com/hero", "https://www.dinamani.com/ta",
                         "https://news.google.com/rss/articles/abc"}
    en = arts["https://www.etvbharat.com/hero"]
    assert en["pub_date"] == "2020-10-14" and en["year"] == 2020 and en["source_country"] == "India"
    ta = arts["https://www.dinamani.com/ta"]
    assert ta["language"] == "Tamil" and "inscription" in json.loads(ta["groups"])
    assert "hero_stone" in json.loads(ta["groups"])
    assert arts["https://news.google.com/rss/articles/abc"]["match_basis"] == "title-only"
    status = {r["url"]: r["status"] for r in db.conn.execute("SELECT url,status FROM candidates")}
    assert status["https://www.heritagedaily.com/kent"] == "rejected"
    assert status["https://dead.example.lk/x"] == "failed"

    out = tmp_path / "out"
    res = export(db, cfg, out)
    assert res["articles"] == 3
    wb = openpyxl.load_workbook(out / "articles.xlsx")
    assert wb.sheetnames == ["Articles", "New (last 7 days)", "By keyword", "By year", "By source", "By place"]
    assert wb["New (last 7 days)"].max_row == 4          # all found today
    assert len(list((out / "daily").glob("*.csv"))) == 1
    ws = wb["Articles"]
    assert ws.max_row == 4 and ws.cell(2, 12).hyperlink is not None
    assert (out / "articles.csv").read_text(encoding="utf-8-sig").count("\n") == 4


def test_country_for(cfg):
    assert verify.country_for("www.dailymirror.lk", cfg) == "Sri Lanka"
    assert verify.country_for("epaper.thehindu.com", cfg) == "India"
    assert verify.country_for("unknown.example.com", cfg) == "Unknown"
    assert verify.country_for("www.tnpscthervupettagam.com", cfg) == "India"
    assert verify.country_for("news.ahram.org.eg", cfg) == "Egypt"


def test_common_noise_rejected(matcher):
    assert matcher.accept("Illegal sand excavation in Tamil Nadu", "Officials seized lorries in Tiruchi.") is None
    assert matcher.accept("Police unearthed a scam", "Arrests in Chennai, Tamil Nadu.") is None
    assert not matcher.match("அவர் தமிழில் பேசினார்").ok


def test_backfill_checks_links_even_when_queue_small(monkeypatch, cfg, tmp_path):
    """Regression: a backfill that finds < 300 links must still open and check them."""
    import collector.__main__ as cli
    from collector.sources import date_archive as da, gdelt as gd, wayback as wb

    monkeypatch.setattr(cli, "DATA_DIR", tmp_path)
    monkeypatch.setattr(cli, "add_seed_urls",
                        lambda db: db.add_candidates([{"url": "https://www.etvbharat.com/hero",
                                                       "collector": "seed"}]))
    calls = {"n": 0}

    def slow_wayback(db, cfg):
        calls["n"] += 1
        db.add_candidates([{"url": f"https://www.dinamani.com/ta{calls['n']}", "collector": "wayback:x"}])
        return calls["n"] < 5

    monkeypatch.setattr(wb, "step", slow_wayback)
    monkeypatch.setattr(da, "step", lambda *a, **k: False)
    monkeypatch.setattr(gd, "step", lambda *a, **k: False)
    pages = {"https://www.etvbharat.com/hero": ARTICLE_EN}
    pages.update({f"https://www.dinamani.com/ta{i}": ARTICLE_TA for i in range(1, 6)})
    monkeypatch.setattr(http, "get", lambda url, **kw: resp(pages.get(url, ""), 200 if url in pages else 404, url))

    db = DB(tmp_path / "c.sqlite")
    cli.run_backfill(db, cfg, Matcher(cfg), minutes=0.2)
    assert db.stats()["matched"] == 6
    assert db.count_pending() == 0
    assert (tmp_path / "status.txt").read_text(encoding="utf-8").startswith("Last updated")
    assert (tmp_path / "articles.xlsx").exists()


def test_run_all_checks_todays_news_first(monkeypatch, cfg, tmp_path):
    import collector.__main__ as cli
    from collector.sources import gdelt as gd, gnews as gn

    monkeypatch.setattr(cli, "DATA_DIR", tmp_path)
    monkeypatch.setattr(cli, "add_seed_urls", lambda db: 0)
    monkeypatch.setattr(gn, "recent", lambda db, cfg, days=7, **k: db.add_candidates(
        [{"url": "https://www.dinamani.com/today", "collector": "gnews", "title": "நடுகல் கண்டுபிடிப்பு"}]))
    monkeypatch.setattr(gd, "recent", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("GDELT down")))
    monkeypatch.setattr(http, "get", lambda url, **kw: resp(ARTICLE_TA, 200, url))
    ran = {}
    monkeypatch.setattr(cli, "run_backfill", lambda db, cfg, m, minutes: ran.setdefault("backfill", minutes))

    db = DB(tmp_path / "r.sqlite")
    cli.run_all(db, cfg, Matcher(cfg), minutes=60)
    assert db.stats()["matched"] == 1                 # today's article checked despite GDELT failure
    assert ran["backfill"] > 5                        # remaining time handed to history collection


def test_search_queries_loaded(cfg):
    assert any("நடுகல்" in q for q in cfg.search_queries_ta)
    assert len(cfg.search_queries_ta) >= 20 and len(cfg.search_queries_en) >= 25


def test_gdelt_recent_gives_up_after_repeated_failures(monkeypatch, cfg, db):
    calls = []

    def failing(*a, **k):
        calls.append(1)
        raise RuntimeError("HTTP 429")

    monkeypatch.setattr(gdelt, "search", failing)
    assert gdelt.recent(db, cfg) == 0
    assert len(calls) == 3                       # not all 57 queries


def test_gnews_respects_time_budget(monkeypatch, cfg):
    from collector.sources import gnews
    calls = []
    monkeypatch.setattr(http, "get", lambda *a, **k: calls.append(1) or resp("<rss></rss>"))
    gnews.fetch(cfg, recent_days=7, budget_min=0)
    assert calls == []                           # budget already used up: no requests


def test_wayback_uses_sections_for_big_sites(cfg):
    jobs = wayback.plan(cfg)
    hindu = [j for j in jobs if j["domain"] == "thehindu.com"]
    assert hindu and all(j["match"] == "prefix" for j in hindu)
    assert any(j["target"] == "thehindu.com/news/national/tamil-nadu/" for j in hindu)
    small = [j for j in jobs if j["domain"] == "dinakaran.com"]
    assert small and all(j["match"] == "domain" and j["target"] == "dinakaran.com" for j in small)
    p = wayback._params(hindu[0], cfg)
    assert p["matchType"] == "prefix" and p["url"].startswith("thehindu.com/news/")


@pytest.mark.parametrize("raw,expected", [
    ("http://www.thehindu.com/news/x/terracotta-pipes-keeladi/article68464783.ece/amp/",
     "https://www.thehindu.com/news/x/terracotta-pipes-keeladi/article68464783.ece"),
    ("https://www.dinamalar.com/news/a/hero-stone/123?utm=1#top", "https://www.dinamalar.com/news/a/hero-stone/123"),
    ("https://www.dtnext.in/news/tamilnadu/hero-stone-found/amp", "https://www.dtnext.in/news/tamilnadu/hero-stone-found/"),
])
def test_canonical(raw, expected):
    assert wayback.canonical(raw) == expected


def test_image_alternates_skipped():
    u = "https://www.thehindu.com/news/cities/Madurai/10m5x4/article32239394.ece/ALTERNATES/LANDSCAPE_615/INSCRIPTION"
    assert wayback.SKIP_PATH.search(u)


def test_live_paywall_falls_back_to_archive(monkeypatch, cfg, db, matcher):
    paywall = "<html><body><article><p>" + ("Subscribe to continue reading. " * 20) + "</p></article></body></html>"

    def fake_get(url, **kw):
        if url.startswith("https://web.archive.org/"):
            return resp(ARTICLE_EN, 200, url)
        return resp(paywall, 200, url)

    monkeypatch.setattr(http, "get", fake_get)
    db.add_candidate("https://www.thehindu.com/hero.ece", "wayback:thehindu.com", wayback_ts="20200101000000")
    db.commit()
    verify.verify_batch(db, cfg, matcher, workers=1)
    a = db.conn.execute("SELECT * FROM articles").fetchone()
    assert a is not None and a["archive_url"].startswith("https://web.archive.org/web/20200101000000/")
